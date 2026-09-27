"""Stage 2 - Extract.

Turn heterogeneous documents (PDF / DOCX / XLSX, possibly 100+ pages) into a
flat list of text chunks that each carry provenance: which document, which
page, which section. Provenance is not decoration -- it is what makes a later
risk judgement checkable by a human, and what lets us reject an LLM citation
that points at a chunk which does not exist.

Chunking strategy: split on natural document boundaries first (page, table,
worksheet), then pack to a target character budget with a small overlap so a
sentence straddling a boundary is not lost to retrieval. Character-based rather
than token-based because it is deterministic, dependency-free, and the
difference does not matter at this granularity.
"""

from __future__ import annotations

import io
from typing import Iterable

from ..models import Chunk, DocumentRef
from ..storage import ObjectStore

# --- per-format readers ----------------------------------------------------


def _read_pdf(data: bytes) -> list[tuple[int | None, str, str]]:
    import pdfplumber

    out: list[tuple[int | None, str, str]] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                out.append((page_no, f"page {page_no}", text))
    return out


def _read_docx(data: bytes) -> list[tuple[int | None, str, str]]:
    import docx

    document = docx.Document(io.BytesIO(data))
    out: list[tuple[int | None, str, str]] = []

    current_heading = "body"
    buffer: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("Heading"):
            if buffer:
                out.append((None, current_heading, "\n".join(buffer)))
                buffer = []
            current_heading = text
        buffer.append(text)
    if buffer:
        out.append((None, current_heading, "\n".join(buffer)))

    for t_idx, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            out.append((None, f"table {t_idx}", "\n".join(rows)))
    return out


def _read_xlsx(data: bytes) -> list[tuple[int | None, str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out: list[tuple[int | None, str, str]] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            out.append((None, f"sheet:{sheet.title}", "\n".join(rows)))
    wb.close()
    return out


_READERS = {"pdf": _read_pdf, "docx": _read_docx, "xlsx": _read_xlsx}


# --- chunking --------------------------------------------------------------


def _pack(text: str, target: int, overlap: int) -> Iterable[str]:
    text = text.strip()
    if len(text) <= target:
        if text:
            yield text
        return

    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    buf = ""
    for para in paragraphs:
        if len(buf) + len(para) + 1 <= target:
            buf = f"{buf}\n{para}" if buf else para
        else:
            if buf:
                yield buf
                buf = (buf[-overlap:] + "\n" + para) if overlap else para
            else:
                # A single oversized paragraph: hard-split it.
                for i in range(0, len(para), target):
                    yield para[i:i + target]
                buf = ""
    if buf.strip():
        yield buf.strip()


def run(documents: list[DocumentRef], store: ObjectStore, logger,
        target_chars: int, overlap_chars: int, skip_doc_ids: set[str] | None = None) -> list[Chunk]:
    skip = skip_doc_ids or set()
    chunks: list[Chunk] = []

    for doc in documents:
        if doc.doc_id in skip:
            continue
        reader = _READERS.get(doc.media_type)
        if reader is None:
            logger.warn("extract.unsupported_media", doc_id=doc.doc_id, media_type=doc.media_type)
            continue
        try:
            segments = reader(store.read_bytes(doc.uri))
        except Exception as exc:  # a corrupt file is a finding, not a crash
            logger.error("extract.failed", doc_id=doc.doc_id, error=repr(exc))
            continue

        made = 0
        for page, section, text in segments:
            for piece in _pack(text, target_chars, overlap_chars):
                chunk_id = f"{doc.doc_id}:{made:03d}"
                chunks.append(
                    Chunk(
                        chunk_id=chunk_id,
                        doc_id=doc.doc_id,
                        filename=doc.filename,
                        page=page,
                        section=section,
                        text=piece,
                    )
                )
                made += 1
        logger.log("extract.document", doc_id=doc.doc_id, media_type=doc.media_type,
                   segments=len(segments), chunks=made)

    logger.log("extract.ok", chunks=len(chunks),
               chars=sum(c.char_count for c in chunks))
    return chunks
